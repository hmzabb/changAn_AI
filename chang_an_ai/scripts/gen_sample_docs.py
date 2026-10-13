"""生成 PDF 和 Word 格式的测试语料（含表格），用于验证多格式入库流程。

用法：
    .venv\\Scripts\\python scripts/gen_sample_docs.py
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.config import BASE_DIR

CORPUS = BASE_DIR / "app" / "data" / "corpus"


def gen_docx_attraction():
    """生成 attractions/西安碑林博物馆.docx —— 含门票价格表格。"""
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.oxml.ns import qn

    doc = Document()

    # 设置默认字体
    style = doc.styles["Normal"]
    font = style.font
    font.name = "宋体"
    font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    doc.add_heading("西安碑林博物馆", level=1)
    doc.add_paragraph(
        "西安碑林博物馆是中国最大的石质书库，位于西安市碑林区三学街15号，"
        "紧邻西安古城墙文昌门。馆藏文物1万余件，其中国宝级文物19组134件，"
        "以汉唐碑刻、历代书法名碑闻名于世，被誉为儒家图书馆和书法艺术宝库。"
    )

    doc.add_heading("馆藏亮点", level=2)
    doc.add_paragraph(
        "碑林始建于北宋元祐二年（1087年），为保存唐代《开成石经》而建，"
        "距今已有九百余年历史。馆内珍藏《大秦景教流行中国碑》、《多宝塔碑》、"
        "《颜氏家庙碑》、《玄秘塔碑》等书法史上的巅峰之作。"
        "第一展室陈列的《开成石经》共114石、65万余字，完整保存了12部儒家经典，"
        "被誉为中国最大最重的一部书。此外，石刻艺术馆收藏了大量汉代至清代的"
        "石刻造像、陵墓石刻和宗教石刻，以昭陵六骏最为著名。"
    )

    doc.add_heading("参观信息", level=2)
    doc.add_paragraph(
        "碑林博物馆全年开放（除夕闭馆一天），建议预留2-3小时参观。"
        "旺季建议上午9点前到达，避开人流高峰。馆内可租赁语音导览器，"
        "也有志愿者免费讲解服务。周边有书院门文化街，参观后可顺路游览。"
    )

    # 表格：门票价格
    doc.add_heading("门票价格一览", level=2)
    table = doc.add_table(rows=6, cols=3, style="Table Grid")
    headers = ["票种", "价格（元）", "适用人群"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True

    data = [
        ["旺季全价票", "65", "成人（3月1日-11月30日）"],
        ["淡季全价票", "50", "成人（12月1日-次年2月底）"],
        ["学生半价票", "32", "全日制学生凭学生证"],
        ["免票", "0", "65岁以上老人、残疾人、现役军人"],
        ["碑林+城墙联票", "100", "成人（仅旺季有售）"],
    ]
    for row_idx, row_data in enumerate(data, 1):
        for col_idx, val in enumerate(row_data):
            table.rows[row_idx].cells[col_idx].text = val

    doc.add_paragraph("")
    doc.add_paragraph("以上票价以现场公告为准，节假日可能有调整。")

    doc.add_heading("交通指南", level=2)
    doc.add_paragraph(
        "地铁：2号线永宁门站D2口出，沿城墙向东步行约10分钟。\n"
        "公交：乘坐14路、23路、40路、208路至文昌门站下车。\n"
        "自驾：导航至碑林博物馆停车场，车位有限，建议公共交通出行。"
    )

    out = CORPUS / "attractions" / "西安碑林博物馆.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"[生成] {out.relative_to(BASE_DIR)}")


def gen_pdf_guide():
    """生成 guides/西安三日游详细行程.pdf —— 含行程安排表格。"""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_font("SimSun", "", r"C:\Windows\Fonts\simsun.ttc", uni=True)
    pdf.add_font("SimHei", "", r"C:\Windows\Fonts\simhei.ttf", uni=True)
    pdf.add_page()

    # 标题
    pdf.set_font("SimHei", "", 18)
    pdf.cell(0, 12, "西安三日游详细行程", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(4)

    # 正文
    pdf.set_font("SimSun", "", 11)
    intro = (
        "西安，古称长安，是十三朝古都，丝绸之路起点。这份三日游行程涵盖了古都核心景区、"
        "美食体验与文化深度游，适合首次来西安的游客。行程张弛有度，既有热门景点打卡，"
        "也安排了小众深度体验，让你在72小时内领略长安的千年风华。"
    )
    pdf.multi_cell(0, 7, intro)
    pdf.ln(4)

    # 行程安排表格
    pdf.set_font("SimHei", "", 14)
    pdf.cell(0, 10, "行程安排总览", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    pdf.set_font("SimSun", "", 9)
    col_w = [38, 38, 38, 76]
    headers = ["时间", "地点", "类型", "详细说明"]
    pdf.set_font("SimHei", "", 9)
    for i, h in enumerate(headers):
        pdf.cell(col_w[i], 8, h, border=1, align="C")
    pdf.ln()

    schedule = [
        ["Day1 上午", "兵马俑", "景点", "世界第八大奇迹，建议8:30抵达，游览约3小时，请导游讲解"],
        ["Day1 中午", "临潼", "美食", "临潼石榴、火晶柿子、老白家肉夹馍"],
        ["Day1 下午", "华清宫", "景点", "骊山脚下皇家园林，长恨歌实景演出（需另购票）"],
        ["Day1 晚上", "回民街", "美食", "羊肉泡馍、灌汤包、红柳烤肉、花奶奶酸梅汤"],
        ["Day2 上午", "西安城墙", "景点", "骑行14公里古城墙（永宁门上最佳），建议9点前上墙"],
        ["Day2 中午", "钟楼附近", "美食", "德发长饺子宴、同盛祥牛羊肉泡馍"],
        ["Day2 下午", "碑林博物馆", "景点", "中国最大石质书库，书法爱好者的天堂"],
        ["Day2 晚上", "大唐不夜城", "休闲", "不倒翁小姐姐、盛唐密盒、灯光秀，建议19:30后到达"],
        ["Day3 上午", "陕西历史博物馆", "景点", "提前3天预约！珍宝馆30元值得进，基本馆免费"],
        ["Day3 中午", "小寨赛格", "美食", "长安大排档、毛笔酥、葫芦鸡"],
        ["Day3 下午", "大雁塔", "景点", "大慈恩寺内，可登塔俯瞰西安。旁边大唐芙蓉园可选"],
        ["Day3 晚上", "洒金桥", "美食", "老米家大雨肉夹馍、马二酸汤水饺、刘信小炒泡馍"],
    ]

    pdf.set_font("SimSun", "", 9)
    for row in schedule:
        for i, cell in enumerate(row):
            pdf.cell(col_w[i], 8, cell, border=1, align="C" if i < 3 else "L")
        pdf.ln()

    pdf.ln(6)

    # 费用总览
    pdf.set_font("SimHei", "", 14)
    pdf.cell(0, 10, "费用预估", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    col_w2 = [80, 50, 60]
    pdf.set_font("SimHei", "", 9)
    for h in ["项目", "预估花费（元/人）", "备注"]:
        pdf.cell(col_w2[h == "项目" and 0 or h == "预估花费（元/人）" and 1 or 2], 8, h, border=1, align="C")
    pdf.ln()

    costs = [
        ["门票总计（以上全部景点）", "约400-500", "学生证可享半价，部分景点需预约"],
        ["餐饮（3天）", "约300-500", "小吃为主，丰俭由人"],
        ["交通（地铁+打车）", "约100-150", "地铁覆盖主要景区"],
        ["住宿（2晚）", "约400-800", "建议住钟楼/南门附近"],
        ["合计", "约1200-1950", "不含往返大交通"],
    ]
    pdf.set_font("SimSun", "", 9)
    for row in costs:
        for i, cell in enumerate(row):
            pdf.cell(col_w2[i], 8, cell, border=1, align="C" if i == 1 else "L")
        pdf.ln()

    pdf.ln(6)
    pdf.set_font("SimSun", "", 10)
    tips = (
        "温馨提示：以上价格为参考价，实际以现场为准。旺季（4-10月）建议提前1-2周预订酒店。"
        "西安夏季炎热，做好防晒；冬季干燥寒冷，注意保暖。"
    )
    pdf.multi_cell(0, 7, tips)

    out = CORPUS / "guides" / "西安三日游详细行程.pdf"
    out.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(out))
    print(f"[生成] {out.relative_to(BASE_DIR)}")


# ==================== 新增：华清宫与骊山 .docx ====================

def gen_docx_huaqing():
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    doc.add_heading("华清宫与骊山", level=1)
    doc.add_paragraph(
        "华清宫位于西安市临潼区骊山北麓，距西安市区约30公里，是中国古代四大皇家园林之一。"
        "这里南依骊山，北临渭水，以温泉汤池闻名，周、秦、汉、隋、唐历代帝王均在此修建离宫别苑。"
        "华清宫不仅有秀美的自然风光，更承载了唐玄宗与杨贵妃的千古爱情传奇。"
    )

    doc.add_heading("历史沿革", level=2)
    doc.add_paragraph(
        "华清宫始建于西周时期，初名骊山汤。唐贞观十八年（644年），唐太宗在此建汤泉宫。"
        "唐玄宗天宝六年（747年）大规模扩建，更名为华清宫。每年十月至次年暮春，"
        "唐玄宗携杨贵妃在此居住理政，诗人白居易在《长恨歌》中写下了春寒赐浴华清池，"
        "温泉水滑洗凝脂的千古名句。1936年，震惊中外的西安事变也发生于此，"
        "五间厅至今保留着当时蒋介石居住的原貌。"
    )

    doc.add_heading("主要景点", level=2)
    doc.add_paragraph(
        "华清宫景区包含御汤遗址博物馆、九龙湖、芙蓉湖、唐梨园遗址等核心景点。"
        "御汤遗址博物馆是国内唯一的皇家汤池遗址专题博物馆，游客可近距离观看"
        "莲花汤（玄宗御用）、海棠汤（贵妃御用）、星辰汤等唐代浴池遗迹。"
        "骊山索道可将游客送上山顶，俯瞰关中平原，参观周幽王烽火戏诸侯的"
        "烽火台遗址和老母殿等古迹。"
    )

    doc.add_heading("演出推荐", level=2)
    doc.add_paragraph(
        "大型实景历史舞剧《长恨歌》每晚在华清宫九龙湖上演，以骊山为背景，"
        "水上舞台配合高科技灯光特效，再现唐玄宗与杨贵妃的爱情故事。"
        "演出时长约70分钟，建议提前购票（旺季一票难求）。"
        "另有《12·12》西安事变实景影画剧，沉浸式再现历史瞬间。"
    )

    doc.add_heading("门票与开放时间", level=2)
    table = doc.add_table(rows=7, cols=3, style="Table Grid")
    for i, h in enumerate(["项目", "价格/时间", "备注"]):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    data = [
        ["旺季门票（3-11月）", "120元/人", "含华清宫景区，不含骊山索道"],
        ["淡季门票（12-2月）", "80元/人", "含华清宫景区，不含骊山索道"],
        ["骊山索道", "往返60元/人", "单程35元/人"],
        ["《长恨歌》演出", "298-988元/人", "按座位区域定价，19:30开演"],
        ["开放时间", "旺季7:00-19:00", "淡季7:30-18:00"],
        ["建议游览时长", "3-4小时", "含骊山则需要半天"],
    ]
    for row_idx, row_data in enumerate(data, 1):
        for col_idx, val in enumerate(row_data):
            table.rows[row_idx].cells[col_idx].text = val

    doc.add_paragraph("以上信息以景区当日公告为准。")

    doc.add_heading("交通指南", level=2)
    doc.add_paragraph(
        "公交：西安火车站东广场乘坐游5（306路）至华清池站，约1小时。\n"
        "地铁：9号线华清池站C口出，步行约5分钟。\n"
        "自驾：连霍高速临潼出口下，导航至华清宫停车场。"
    )

    out = CORPUS / "attractions" / "华清宫与骊山.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"[生成] {out.relative_to(BASE_DIR)}")


# ==================== 新增：西安面食地图 .docx ====================

def gen_docx_noodles():
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    doc.add_heading("西安面食地图", level=1)
    doc.add_paragraph(
        "西安是当之无愧的中国面食之都。关中平原盛产优质小麦，八百里秦川养成了"
        "西安人对面食的极致追求。从街头小吃到宴席大菜，一碗面可以变幻出上百种花样。"
        "以下带你走进西安面食江湖，了解每一碗面背后的故事与讲究。"
    )

    doc.add_heading("面条江湖：十大经典面食", level=2)
    doc.add_paragraph(
        "西安面食的核心在于筋道。陕西小麦蛋白质含量高，手工和面、反复揉搓、充分醒发，"
        "才能拉出薄如纸、细如发的面条。关中八大怪中就有面条像裤带的说法，"
        "说的正是宽如腰带的Biangbiang面。吃面讲究油泼辣子一道菜，"
        "秦椒碾制的辣椒面浇上滚烫的热油，滋啦一声香气四溢，是西安面食的灵魂所在。"
    )

    doc.add_heading("经典面食一览", level=2)
    table = doc.add_table(rows=11, cols=4, style="Table Grid")
    for i, h in enumerate(["面食名称", "参考价格", "口味特点", "推荐去处"]):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    data = [
        ["Biangbiang面", "15-25元", "宽如裤带，筋道弹牙，油泼辣子+臊子浇头", "南院门Biangbiang面"],
        ["油泼面", "12-18元", "滚油浇辣椒面，蒜香浓郁，简单纯粹", "柳巷面馆、老白家"],
        ["岐山臊子面", "15-22元", "酸辣鲜香，薄筋光煎稀旺，臊子五彩缤纷", "永丰岐山面"],
        ["羊肉泡馍", "25-45元", "汤浓肉烂，亲手掰馍黄豆粒大小最佳", "同盛祥、老孙家"],
        ["肉夹馍（腊汁肉）", "10-18元", "白吉馍酥脆，腊汁肉肥而不腻", "秦豫肉夹馍、子午路张记"],
        ["葫芦头", "25-38元", "猪大肠+馍，汤白味醇，配糖蒜解腻", "春发生、朱秀英"],
        ["麻酱凉皮", "8-15元", "酸辣爽滑，芝麻酱浓郁，夏日消暑利器", "盛志望麻酱酿皮"],
        ["灌汤包", "18-28元", "皮薄汤鲜，先开窗后喝汤", "贾三灌汤包子"],
        ["酸汤水饺", "15-22元", "牛肉韭黄馅，酸汤开胃，油泼辣子增香", "马二酸汤水饺"],
        ["水盆羊肉", "28-40元", "汤清肉烂，配月牙饼同食", "澄城崇斌水盆"],
    ]
    for row_idx, row_data in enumerate(data, 1):
        for col_idx, val in enumerate(row_data):
            table.rows[row_idx].cells[col_idx].text = val

    doc.add_paragraph("")
    doc.add_paragraph("以上价格为2025年参考价，实际以店铺菜单为准。")

    doc.add_heading("吃面学问", level=2)
    doc.add_paragraph(
        "西安面馆的桌上标配：油泼辣子、醋、生蒜。一碗面端上来，先别急着拌，"
        "陕西人讲究面要趁热吃，坨了就失了魂。桌上生蒜剥两瓣，一口面一口蒜，"
        "这才是地道吃法。吃泡馍有讲究：掰馍要耐心，黄豆粒大小最入味，"
        "掰完交给师傅煮馍，单走（汤馍分离）还是口汤（收汁刚好）全凭个人喜好。"
    )

    out = CORPUS / "food" / "西安面食地图.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"[生成] {out.relative_to(BASE_DIR)}")


# ==================== 新增：西安交通与住宿指南 .docx ====================

def gen_docx_transport():
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    doc.add_heading("西安交通与住宿指南", level=1)
    doc.add_paragraph(
        "西安作为西北地区交通枢纽，拥有完善的航空、高铁、地铁和公交网络。"
        "无论你是飞机、高铁还是自驾抵达，这份指南都能帮你迅速规划出行路线。"
    )

    doc.add_heading("抵达西安", level=2)
    doc.add_paragraph(
        "西安咸阳国际机场位于市区西北方向约40公里，目前有T1、T2、T3三个航站楼。"
        "机场至市区可乘坐地铁14号线（约50分钟至北客站），或机场大巴（25元/人，"
        "多条线路覆盖市区主要区域）。西安北站是西北最大的高铁枢纽，郑西、西成、"
        "大西、银西等多条高铁交汇于此。西安站位于市中心城墙北侧，主要接发普速列车。"
    )

    doc.add_heading("市内交通方式对比", level=2)
    table = doc.add_table(rows=7, cols=4, style="Table Grid")
    for i, h in enumerate(["交通方式", "票价", "运营时间", "适用场景"]):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    data = [
        ["地铁", "2-8元/次", "6:00-23:30", "覆盖主要景区，准时快捷，高峰期拥挤"],
        ["公交车", "1-2元/次", "6:00-22:00", "线路密集，覆盖地铁盲区，速度较慢"],
        ["出租车/网约车", "起步价8.5元", "全天候", "行李多、赶时间时首选，高峰期难打车"],
        ["共享单车", "1.5元/30分钟", "全天候", "城墙内短途接驳，春秋季骑游城墙周边最佳"],
        ["旅游专线", "7-10元/次", "7:00-18:00", "直达兵马俑、华清宫等远郊景点"],
        ["景区直通车", "25-35元", "8:00-16:00", "酒店直达景区，含导游讲解"],
    ]
    for row_idx, row_data in enumerate(data, 1):
        for col_idx, val in enumerate(row_data):
            table.rows[row_idx].cells[col_idx].text = val

    doc.add_heading("住宿区域推荐", level=2)
    doc.add_paragraph(
        "钟楼/南门区域：西安最核心地段，步行可至回民街、城墙，酒店选择丰富，"
        "从青旅到五星级均有，价格200-1500元不等。大雁塔/曲江区域：环境优美，"
        "高档酒店集中，夜游大唐不夜城方便，适合追求品质的游客。小寨/体育场区域："
        "商业繁华，购物方便，地铁2/3号线交汇，性价比高。火车站/北关区域："
        "经济型酒店集中，适合预算有限的背包客，但夜间环境相对嘈杂。"
    )

    doc.add_heading("住宿预算对比", level=2)
    table2 = doc.add_table(rows=5, cols=3, style="Table Grid")
    for i, h in enumerate(["档次", "价格区间（元/晚）", "推荐"]):
        cell = table2.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    data2 = [
        ["经济型/青旅", "50-200", "湘子门青旅、西安七贤青旅"],
        ["舒适型/连锁", "200-500", "全季、亚朵、美居"],
        ["高档型", "500-1000", "西安威斯汀、凯悦、W酒店"],
        ["奢华型", "1000以上", "西安索菲特传奇、丽思卡尔顿"],
    ]
    for row_idx, row_data in enumerate(data2, 1):
        for col_idx, val in enumerate(row_data):
            table2.rows[row_idx].cells[col_idx].text = val

    out = CORPUS / "guides" / "西安交通与住宿指南.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"[生成] {out.relative_to(BASE_DIR)}")


# ==================== 新增：兵马俑参观指南 .pdf ====================

def gen_pdf_terracotta():
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_font("SimSun", "", r"C:\Windows\Fonts\simsun.ttc", uni=True)
    pdf.add_font("SimHei", "", r"C:\Windows\Fonts\simhei.ttf", uni=True)
    pdf.add_page()

    pdf.set_font("SimHei", "", 18)
    pdf.cell(0, 12, "兵马俑参观指南", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(4)

    pdf.set_font("SimSun", "", 10)
    intro = (
        "秦始皇兵马俑博物馆位于西安市临潼区，距市区约35公里，是世界第八大奇迹，"
        "联合国教科文组织世界文化遗产。1974年当地农民打井时偶然发现，随后考古发掘"
        "震惊世界。目前已发现四个俑坑，总面积超过2万平方米，出土陶俑、陶马8000余件，"
        "青铜兵器数万件，再现了秦代军队的磅礴气势。"
    )
    pdf.multi_cell(0, 6, intro)
    pdf.ln(3)

    pdf.set_font("SimHei", "", 13)
    pdf.cell(0, 9, "各坑参观指南", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    col_w = [30, 42, 58, 60]
    pdf.set_font("SimHei", "", 8)
    for h in ["坑号", "面积/规模", "主要看点", "参观建议"]:
        pdf.cell(col_w[["坑号", "面积/规模", "主要看点", "参观建议"].index(h)], 7, h, border=1, align="C")
    pdf.ln()

    pits = [
        ["一号坑", "14260㎡，6000+陶俑", "主力军阵，步兵+战车，气势最震撼", "最先参观，建议请导游讲解军阵布局"],
        ["二号坑", "6000㎡，1300+陶俑", "多兵种混合，骑兵+弩兵+战车", "兵马俑中的精华，将军俑在此出土"],
        ["三号坑", "520㎡，68尊陶俑", "指挥部/军幕，级别最高", "面积小但地位最高，注意观察布局"],
        ["四号坑", "未建成，空坑", "推测因秦末战乱未完工", "了解秦朝覆灭的历史背景"],
    ]
    pdf.set_font("SimSun", "", 8)
    for row in pits:
        for i, cell in enumerate(row):
            pdf.cell(col_w[i], 7, cell, border=1, align="C" if i < 2 else "L")
        pdf.ln()

    pdf.ln(4)

    pdf.set_font("SimHei", "", 13)
    pdf.cell(0, 9, "门票与实用信息", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    col_w2 = [70, 65, 55]
    pdf.set_font("SimHei", "", 8)
    for h in ["项目", "详情", "备注"]:
        pdf.cell(col_w2[["项目", "详情", "备注"].index(h)], 7, h, border=1, align="C")
    pdf.ln()

    info = [
        ["旺季门票（3-11月）", "120元/人", "学生半价，65岁以上免票"],
        ["淡季门票（12-2月）", "120元/人", "兵马俑全年同价"],
        ["开放时间", "8:30-17:00（旺季至17:30）", "建议早8:30前到，避开人流高峰"],
        ["讲解服务", "导游150元/次，讲解器30元/台", "强烈建议请讲解，否则走马观花"],
        ["建议游览时长", "2.5-3.5小时", "含秦始皇陵（联票已含摆渡车）"],
        ["交通方式", "游5(306路)/地铁9号线", "火车站东广场乘游5，约1小时"],
    ]
    pdf.set_font("SimSun", "", 8)
    for row in info:
        for i, cell in enumerate(row):
            pdf.cell(col_w2[i], 7, cell, border=1, align="C" if i == 0 else "L")
        pdf.ln()

    pdf.ln(4)
    pdf.set_font("SimHei", "", 13)
    pdf.cell(0, 9, "参观路线推荐", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    pdf.set_font("SimSun", "", 10)
    route = (
        "推荐参观顺序：一号坑（震撼开场）→ 三号坑（指挥中枢）→ 二号坑（精华所在）"
        "→ 文物陈列厅（铜车马必看）。全程约2.5小时。参观完毕后乘坐免费摆渡车前往"
        "秦始皇陵（丽山园），车程约10分钟，游览约40分钟。返回时可顺路参观华清宫。\n\n"
        "拍照提示：一号坑正面观景台人最多，可走到两侧人少处拍摄；坑内禁用闪光灯。"
        "二号坑的跪射俑是兵马俑博物馆的镇馆之宝，保存最完整、色彩最丰富。"
    )
    pdf.multi_cell(0, 6, route)

    out = CORPUS / "attractions" / "兵马俑参观指南.pdf"
    out.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(out))
    print(f"[生成] {out.relative_to(BASE_DIR)}")


# ==================== 新增：回民街美食推荐 .pdf ====================

def gen_pdf_food_street():
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_font("SimSun", "", r"C:\Windows\Fonts\simsun.ttc", uni=True)
    pdf.add_font("SimHei", "", r"C:\Windows\Fonts\simhei.ttf", uni=True)
    pdf.add_page()

    pdf.set_font("SimHei", "", 18)
    pdf.cell(0, 12, "回民街美食推荐", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(4)

    pdf.set_font("SimSun", "", 10)
    intro = (
        "回民街是西安最著名的美食街区，广义上包含北院门、西羊市、大皮院、洒金桥等"
        "多条街巷。这里是西安回民聚居区，以清真美食闻名，已有上千年历史。"
        "初次到访的游客往往只逛北院门主街，但真正的老饕都知道，美食精华藏在"
        "西羊市和洒金桥的深巷里。以下为你解锁回民街的正确打开方式。"
    )
    pdf.multi_cell(0, 6, intro)
    pdf.ln(3)

    pdf.set_font("SimHei", "", 13)
    pdf.cell(0, 9, "必吃美食清单（含价格参考）", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    col_w = [37, 30, 55, 68]
    pdf.set_font("SimHei", "", 8)
    for h in ["美食名称", "参考价格（元）", "口味特点", "推荐店铺"]:
        pdf.cell(col_w[["美食名称", "参考价格（元）", "口味特点", "推荐店铺"].index(h)], 7, h, border=1, align="C")
    pdf.ln()

    foods = [
        ["羊肉泡馍", "28-45", "汤浓肉烂，掰馍黄豆粒大小，配糖蒜辣酱", "老米家大雨肉夹馍（西羊市）"],
        ["腊牛肉夹馍", "15-20", "肉烂汁浓，白吉馍外酥里软", "杨天玉腊牛肉夹馍（洒金桥）"],
        ["灌汤包子", "18-28", "皮薄如纸，汤汁鲜美，先开窗后喝汤", "贾三灌汤包子（北院门）"],
        ["红柳烤肉", "10-15/串", "红柳枝串羊肉，炭火烤制，孜然飘香", "任意街边摊均可，认准现切现串"],
        ["定家小酥肉", "25-30", "牛肉挂糊油炸后蒸制，软烂入味", "定家小酥肉（大皮院）"],
        ["花奶奶酸梅汤", "6-10", "冰镇酸甜解腻，回民街饮品头牌", "花奶奶酸梅汤（大皮院）"],
        ["东南亚甑糕", "10-15", "糯米+红枣+芸豆层层蒸制，甜糯可口", "东南亚甑糕（西羊市）"],
        ["柿子饼", "3-5/个", "临潼火晶柿子+面粉油炸，香甜软糯", "伊古斋黄桂柿子饼"],
        ["卤汁凉粉", "10-15", "凉粉浇热卤汁，芥末蒜汁芝麻酱", "何老二卤汁凉粉（西羊市）"],
        ["刘纪孝腊牛羊肉", "80-120/斤", "逢年过节排队2小时起，西安伴手礼之王", "刘纪孝腊牛羊肉（西羊市口）"],
    ]
    pdf.set_font("SimSun", "", 8)
    for row in foods:
        for i, cell in enumerate(row):
            pdf.cell(col_w[i], 7, cell, border=1, align="C" if i < 2 else "L")
        pdf.ln()

    pdf.ln(4)

    pdf.set_font("SimHei", "", 13)
    pdf.cell(0, 9, "回民街游览攻略", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    pdf.set_font("SimSun", "", 10)
    tips = (
        "最佳时间：下午5点后陆续出摊，晚上7-9点最热闹。建议空着肚子来！\n\n"
        "路线建议：从洒金桥地铁站B口出 → 洒金桥（吃杨天玉肉夹馍+马二酸汤水饺）→ "
        "大皮院（定家小酥肉+花奶奶酸梅汤）→ 西羊市（老米家泡馍+东南亚甑糕）→ "
        "北院门主街（拍照打卡）。全程步行约2小时。\n\n"
        "重要提示：回民街是回民聚居区，请尊重民族习俗，不要携带非清真食品进入。"
        "不要对回民称猪这个字。拍照时注意不要拍到做礼拜的穆斯林。"
    )
    pdf.multi_cell(0, 6, tips)

    out = CORPUS / "food" / "回民街美食推荐.pdf"
    out.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(out))
    print(f"[生成] {out.relative_to(BASE_DIR)}")


# ==================== 新增：西安景点门票与开放时间总览 .pdf ====================

def gen_pdf_ticket_overview():
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_font("SimSun", "", r"C:\Windows\Fonts\simsun.ttc", uni=True)
    pdf.add_font("SimHei", "", r"C:\Windows\Fonts\simhei.ttf", uni=True)
    pdf.add_page()

    pdf.set_font("SimHei", "", 18)
    pdf.cell(0, 12, "西安景点门票与开放时间总览", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(4)

    pdf.set_font("SimSun", "", 10)
    intro = (
        "本文汇总西安主要景点的最新门票价格与开放时间，数据来源于各景区官方公告"
        "及实地调研（2025年版）。票价分为旺季（3月1日-11月30日）和淡季"
        "（12月1日-次年2月底），部分景区全年统一票价。学生、老人、军人等"
        "优惠政策请以景区当日公告为准。"
    )
    pdf.multi_cell(0, 6, intro)
    pdf.ln(3)

    pdf.set_font("SimHei", "", 13)
    pdf.cell(0, 9, "西安市内主要景点", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    col_w = [48, 28, 28, 40, 46]
    pdf.set_font("SimHei", "", 8)
    for h in ["景点名称", "旺季票价", "淡季票价", "开放时间", "建议游览时长"]:
        pdf.cell(col_w[["景点名称", "旺季票价", "淡季票价", "开放时间", "建议游览时长"].index(h)], 7, h, border=1, align="C")
    pdf.ln()

    attractions = [
        ["兵马俑", "120元", "120元", "8:30-17:00", "3小时"],
        ["华清宫", "120元", "80元", "7:00-19:00", "3-4小时"],
        ["西安城墙", "54元", "54元", "8:00-22:00", "2-3小时"],
        ["碑林博物馆", "65元", "50元", "8:00-18:00", "2-3小时"],
        ["陕西历史博物馆", "免费（预约）", "免费（预约）", "9:00-17:30", "2-4小时"],
        ["大雁塔（大慈恩寺）", "40元+25元", "40元+25元", "8:00-18:00", "1-2小时"],
        ["大唐芙蓉园", "120元", "90元", "9:00-22:00", "3-4小时"],
        ["大明宫遗址公园", "60元", "48元", "8:30-18:00", "2-3小时"],
    ]
    pdf.set_font("SimSun", "", 8)
    for row in attractions:
        for i, cell in enumerate(row):
            pdf.cell(col_w[i], 7, cell, border=1, align="C" if i > 0 else "L")
        pdf.ln()

    pdf.ln(4)

    pdf.set_font("SimHei", "", 13)
    pdf.cell(0, 9, "远郊及周边景点", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    col_w2 = [48, 30, 30, 40, 42]
    pdf.set_font("SimHei", "", 8)
    for h in ["景点名称", "票价", "淡季票价", "开放时间", "建议游览时长"]:
        pdf.cell(col_w2[["景点名称", "票价", "淡季票价", "开放时间", "建议游览时长"].index(h)], 7, h, border=1, align="C")
    pdf.ln()

    outskirts = [
        ["华山", "160元", "100元", "全天（索道7:00-19:00）", "1-2天"],
        ["壶口瀑布（陕西侧）", "90元", "90元", "7:00-18:30", "2-3小时"],
        ["法门寺", "120元", "90元", "8:00-17:30", "2-3小时"],
        ["乾陵", "100元", "80元", "8:00-18:00", "2-3小时"],
        ["汉阳陵", "70元", "55元", "8:30-17:00", "2小时"],
    ]
    pdf.set_font("SimSun", "", 8)
    for row in outskirts:
        for i, cell in enumerate(row):
            pdf.cell(col_w2[i], 7, cell, border=1, align="C" if i > 0 and i < 4 else "L")
        pdf.ln()

    pdf.ln(4)
    pdf.set_font("SimHei", "", 13)
    pdf.cell(0, 9, "省钱攻略", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    pdf.set_font("SimSun", "", 10)
    save_tips = (
        "1. 陕历博基本馆免费但需提前3天在公众号预约，旺季名额秒空。\n"
        "2. 持学生证可享大部分景点半价优惠（研究生证部分景区不认，请提前确认）。\n"
        "3. 西安城墙每月15日为开放日，西安市民免费（需预约）。\n"
        "4. 西安旅游年卡（99元/年）包含城墙、碑林等20+景点，适合深度游。\n"
        "5. 多家景区推出联票优惠（如兵马俑+华清宫、碑林+城墙等），可节省20-30%。\n\n"
        "重要提示：以上所有票价和时间为2025年采集，出行前请务必在"
        "各景区官方公众号确认最新信息，节假日可能有临时调整。"
    )
    pdf.multi_cell(0, 6, save_tips)

    out = CORPUS / "guides" / "西安景点门票与开放时间总览.pdf"
    out.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(out))
    print(f"[生成] {out.relative_to(BASE_DIR)}")


def main():
    print("开始生成测试语料...")
    gen_docx_attraction()
    gen_pdf_guide()

    # 新增 3 个 Word + 3 个 PDF
    gen_docx_huaqing()
    gen_docx_noodles()
    gen_docx_transport()
    gen_pdf_terracotta()
    gen_pdf_food_street()
    gen_pdf_ticket_overview()

    print("\n完成！请在 corpus/ 目录下查看生成的文件。")
    print("运行 scripts/ingest.py 即可测试多格式入库。")

if __name__ == "__main__":
    main()